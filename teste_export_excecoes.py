"""
teste_export_excecoes.py
========================
Gera a planilha de revisao de excecoes e confere, recarregando com openpyxl,
congelamento, filtro, formatos, larguras, abas e impressao.

Rodar com pytest (`pytest teste_export_excecoes.py`) ou direto (`python teste_export_excecoes.py`).
"""

from __future__ import annotations

import io
from datetime import date

import pandas as pd
from openpyxl import load_workbook

from pipeline import exportar_fila_excecoes_xlsx, nome_aba_seguro

OPCOES = ["", "Falta injustificada", "Falta justificada", "Erro de relógio — abonar", "Atestado médico"]
COLUNAS = ["nome_colaborador", "razao_social", "data", "situacao", "motivo_validacao", "classificacao", "observacao"]


def _fila() -> pd.DataFrame:
    return pd.DataFrame([
        ["Zélia Ñandú", "Empresa B & Cia/Filial: [X]?", "2026-09-10", "Provável falta real", "Sem batidas " * 20, "", ""],
        ["Ana Souza", "Empresa A", "2026-09-12", "Provável erro de relógio", "Saída sem retorno", "Erro de relógio — abonar", "=1+1"],
        ["Ana Souza", "Empresa A", "2026-09-02", "Intervalo curto", "Almoço de 10 min\x07", "", "obs longa " * 30],
    ], columns=COLUNAS).assign(data=lambda d: pd.to_datetime(d["data"]))


def _carregar(fila: pd.DataFrame):
    dados = exportar_fila_excecoes_xlsx(fila, "09/2026", OPCOES, "01/09/2026 a 30/09/2026")
    return load_workbook(io.BytesIO(dados))


def test_layout_e_formatos():
    wb = _carregar(_fila())
    assert wb.sheetnames == ["Resumo", "Exceções"]
    ws = wb["Exceções"]
    assert [c.value for c in ws[1]] == ["Colaborador", "Empresa", "Data", "Situação", "Motivo", "Decisão", "Status", "Observação"]
    assert ws.freeze_panes == "B2"
    assert ws.auto_filter.ref == "A1:H4"
    assert ws["A1"].fill.fgColor.rgb.endswith("1B4D3E") and ws["A1"].font.bold
    assert ws.row_dimensions[1].height == 30
    # ordenado por empresa > colaborador > data
    assert [(ws.cell(r, 2).value, ws.cell(r, 1).value) for r in (2, 3, 4)][0][0] == "Empresa A"
    assert ws["C2"].value.date() == date(2026, 9, 2) and ws["C3"].value.date() == date(2026, 9, 12)
    assert ws["C2"].number_format == "DD/MM/YYYY"
    assert [ws.cell(r, 7).value for r in (2, 3, 4)] == ["Pendente", "Classificada", "Pendente"]
    assert ws["H3"].value == "=1+1" and ws["H3"].data_type == "s"  # texto, nao formula
    assert "\x07" not in ws["E2"].value
    assert ws.column_dimensions["E"].width >= 45 and ws.column_dimensions["E"].width <= 70
    assert ws.column_dimensions["H"].width >= 40
    assert ws.page_setup.orientation == "landscape" and ws.page_setup.fitToWidth == 1
    assert ws.print_title_rows == "$1:$1"
    assert "&P" in ws.oddFooter.center.text
    assert sum(len(cf.rules) for cf in ws.conditional_formatting) == 4  # azul, vermelho, pendente, zebra


def test_celulas_separadas_com_textos_dificeis():
    motivo = "possível falta parcial de ~4h20 — Espelho de Ponto: Provável falta parcial — confirmar"
    fila = pd.DataFrame([
        ["ALAN DA SILVA FERREIRA", "MAÇANEIRO E GONZAGA LTDA", "2026-09-22", "Provável erro de relógio",
         f'{motivo}, "batidas": 07:31, 13:35' + chr(10) + 'linha 2', "",'disse "ok", e saiu'],
    ], columns=COLUNAS).assign(data=lambda d: pd.to_datetime(d["data"]))
    ws = _carregar(fila)["Exceções"]
    assert [c.value for c in ws[1]] == ["Colaborador", "Empresa", "Data", "Situação", "Motivo", "Decisão", "Status", "Observação"]
    assert ws["A2"].value == "ALAN DA SILVA FERREIRA" and ws["B2"].value == "MAÇANEIRO E GONZAGA LTDA"
    assert ws["E2"].value.startswith(motivo) and chr(10) in ws["E2"].value
    assert ws["H2"].value == 'disse "ok", e saiu'
    # nenhuma linha inteira colada numa so celula
    assert all(not (isinstance(c.value, str) and c.value.count(",") >= 6) for row in ws.iter_rows() for c in row)
    assert ws["A2"].value.count(",") == 0


def test_exportacao_filtrada():
    fila = pd.DataFrame([
        ["Ana", "Empresa A", "2026-09-02", "Intervalo curto", "m1", "", ""],
        ["Bia", "Empresa B", "2026-09-03", "Provável falta real", "m2", "Falta justificada", ""],
        ["Caio", "Empresa B", "2026-09-04", "Provável falta real", "m3", "", ""],
    ], columns=COLUNAS).assign(data=lambda d: pd.to_datetime(d["data"]))
    recorte = fila[fila["razao_social"] == "Empresa B"]
    dados = exportar_fila_excecoes_xlsx(recorte, "09/2026", OPCOES, "set", "Empresa: Empresa B", len(fila))
    wb = load_workbook(io.BytesIO(dados))
    ws = wb["Exceções"]
    assert [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)] == ["Bia", "Caio"]
    assert ws.auto_filter.ref == "A1:H3"
    res = wb["Resumo"]
    assert res["A3"].value == "Filtros: Empresa: Empresa B · Exibindo 2 de 3 exceções"
    valores = {res.cell(r, 1).value: res.cell(r, 2).value for r in range(1, res.max_row + 1)}
    assert valores["Total de exceções"] == 2 and valores["Classificadas"] == 1 and valores["Pendentes"] == 1
    assert valores["Falta justificada"] == 1 and "Empresa A" not in valores
    vazio = load_workbook(io.BytesIO(exportar_fila_excecoes_xlsx(fila.iloc[0:0], "09/2026", OPCOES, "set", "Empresa: X", 3)))
    assert vazio["Exceções"].max_row == 1 and vazio["Resumo"]["A3"].value.endswith("Exibindo 0 de 3 exceções")


def test_resumo():
    ws = _carregar(_fila())["Resumo"]
    valores = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(1, ws.max_row + 1)}
    assert valores["Total de exceções"] == 3 and valores["Classificadas"] == 1 and valores["Pendentes"] == 2
    assert valores["Erro de relógio — abonar"] == 1 and valores["Falta injustificada"] == 0


def test_fila_vazia():
    ws = _carregar(_fila().iloc[0:0])["Exceções"]
    assert ws.max_row == 1 and ws.freeze_panes == "B2" and ws.auto_filter.ref == "A1:H1"


def test_nome_aba_seguro():
    n = nome_aba_seguro("Empresa B & Cia/Filial: [X]? " + "x" * 40, ["Resumo"])
    assert len(n) <= 31 and not any(ch in n for ch in "[]:*?/\\")
    assert nome_aba_seguro("resumo", ["Resumo"]).lower() != "resumo"


if __name__ == "__main__":
    for nome, fn in list(globals().items()):
        if nome.startswith("test_"):
            fn()
            print("ok", nome)
    sai = exportar_fila_excecoes_xlsx(_fila(), "09/2026", OPCOES, "01/09/2026 a 30/09/2026")
    open("amostra_revisao_excecoes.xlsx", "wb").write(sai)
