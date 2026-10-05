"""
teste_banco_dsr.py
==================
Testes do banco de horas (colunas `banco_horas` / `registra_ponto` do cadastro, roteamento
das horas do relogio para o banco, checagens do preflight, leitura de saldo em hh:mm) e do
desconto de DSR (8794) em dias ou em valor.

Rodar com pytest (`pytest teste_banco_dsr.py`) ou direto (`python teste_banco_dsr.py`).
Neste modo, se a pasta "Dados e Docs Reais" existir (ela e' ignorada pelo git), roda tambem
o cenario de 09/2026 com os arquivos reais e imprime o antes/depois.
"""

from __future__ import annotations

import tempfile
from datetime import date, datetime, time, timedelta
from pathlib import Path

import pandas as pd

import pipeline
from banco_horas import ler_banco_horas_adriano, ler_banco_horas_secullum
from pipeline import (
    avisos_banco_horas, carregar_cadastro_colaboradores, exportar_por_empresa,
    executar_checagens_preflight, rodada_1_gerar_fila_de_excecoes, rodada_2_gerar_relacao_de_valores,
)
from validacao import ArquivoInvalidoError, converter_horas_decimais, converter_valor_monetario

PIS_BANCO = "10000000001"    # no banco de horas, usa o relogio
PIS_COLEGA = "10000000002"   # folha normal
PIS_SEM_PONTO = "10000000003"  # no banco de horas, NAO usa o relogio (caso "Adriano")
COLUNAS_CADASTRO = [
    "pis", "matricula", "nome", "codigo_empresa", "razao_social",
    "entrada", "saida_almoco", "retorno_almoco", "saida", "dias_trabalho",
]


def _pasta() -> Path:
    return Path(tempfile.mkdtemp(prefix="banco_dsr_teste_"))


def _linha_cadastro(pis: str, matricula: str, nome: str, **extras) -> dict:
    return {
        "pis": pis, "matricula": matricula, "nome": nome, "codigo_empresa": "01", "razao_social": "EMPRESA TESTE",
        "entrada": "08:00", "saida_almoco": "12:00", "retorno_almoco": "13:00", "saida": "17:00",
        "dias_trabalho": "0,1,2,3,4", **extras,
    }


def _cadastro_em_arquivo(linhas: list[dict], nome: str = "cadastro_teste.xlsx") -> str:
    caminho = _pasta() / nome
    pd.DataFrame(linhas).to_excel(caminho, index=False)
    return str(caminho)


# ---------------------------------------------------------------------------
# Colunas opcionais do cadastro
# ---------------------------------------------------------------------------
def test_banco_horas_aceita_as_variantes_de_sim_e_nao():
    valores = ["sim", "Não", "nao", "S", "n", 1, 0, "TRUE", "false", "x", None]
    esperado = [True, False, False, True, False, True, False, True, False, True, False]  # vazio = padrao (nao)
    linhas = [_linha_cadastro(f"1000000{i:04d}", str(i), f"PESSOA {i}", banco_horas=v) for i, v in enumerate(valores)]
    cadastro = carregar_cadastro_colaboradores(_cadastro_em_arquivo(linhas))
    assert cadastro["banco_horas"].tolist() == esperado
    assert cadastro["banco_horas"].dtype == bool


def test_registra_ponto_padrao_e_sim_e_variantes():
    valores = ["nao", "NÃO", "sim", None]
    linhas = [_linha_cadastro(f"1000000{i:04d}", str(i), f"PESSOA {i}", registra_ponto=v) for i, v in enumerate(valores)]
    cadastro = carregar_cadastro_colaboradores(_cadastro_em_arquivo(linhas))
    assert cadastro["registra_ponto"].tolist() == [False, False, True, True]  # vazio = padrao (sim)


def test_cadastro_antigo_sem_as_colunas_continua_valendo():
    cadastro = carregar_cadastro_colaboradores(_cadastro_em_arquivo([_linha_cadastro(PIS_BANCO, "1", "ANTIGO")]))
    assert cadastro["banco_horas"].dtype == bool and not cadastro["banco_horas"].any()
    assert cadastro["registra_ponto"].dtype == bool and cadastro["registra_ponto"].all()
    assert "valor_dia_dsr" not in cadastro.columns
    assert not {"banco_horas", "registra_ponto", "valor_dia_dsr"} & set(pipeline.COLUNAS_OBRIGATORIAS_CADASTRO)


def test_valor_invalido_aponta_arquivo_linha_e_coluna():
    linhas = [
        _linha_cadastro("10000000001", "1", "A", banco_horas="sim"),
        _linha_cadastro("10000000002", "2", "B", banco_horas="nao"),
        _linha_cadastro("10000000003", "3", "C", banco_horas="talvez"),
    ]
    caminho = _cadastro_em_arquivo(linhas, "cadastro_com_erro.xlsx")
    try:
        carregar_cadastro_colaboradores(caminho)
    except ArquivoInvalidoError as e:
        assert "cadastro_com_erro.xlsx" in str(e) and "linha 4" in str(e) and "banco_horas" in str(e), str(e)
    else:
        raise AssertionError("deveria recusar 'talvez'")


def test_feriados_usa_o_parser_sim_nao_compartilhado():
    import feriados
    import validacao

    assert feriados.ler_sim_nao is validacao.ler_sim_nao
    assert not hasattr(feriados, "_SIM") and not hasattr(feriados, "_NAO")


def test_valor_dia_dsr_aceita_numero_e_texto_brasileiro():
    valores = [62.5, "62,50", "R$ 62,50", "R$ 1.234,56", None]
    linhas = [_linha_cadastro(f"1000000{i:04d}", str(i), f"P{i}", valor_dia_dsr=v) for i, v in enumerate(valores)]
    cadastro = carregar_cadastro_colaboradores(_cadastro_em_arquivo(linhas))
    esperado = [62.5, 62.5, 62.5, 1234.56]
    assert cadastro["valor_dia_dsr"].iloc[:4].tolist() == esperado
    assert pd.isna(cadastro["valor_dia_dsr"].iloc[4])

    caminho = _cadastro_em_arquivo([_linha_cadastro(PIS_BANCO, "1", "A", valor_dia_dsr="abc")], "cad_valor_ruim.xlsx")
    try:
        carregar_cadastro_colaboradores(caminho)
    except ArquivoInvalidoError as e:
        assert "cad_valor_ruim.xlsx" in str(e) and "linha 2" in str(e) and "valor_dia_dsr" in str(e), str(e)
    else:
        raise AssertionError("deveria recusar 'abc'")


# ---------------------------------------------------------------------------
# Conversao de horas (saldo do banco)
# ---------------------------------------------------------------------------
def _horas(valor) -> float:
    return converter_horas_decimais(valor, "arquivo teste, linha 2", "saldo")


def test_converter_horas_decimal_hhmm_negativo_e_celulas_de_hora():
    assert _horas(5.5) == 5.5 and _horas(3) == 3.0
    assert _horas("5,5") == 5.5 and _horas("5.5") == 5.5 and _horas("1.234,5") == 1234.5
    assert _horas("05:30") == 5.5 and _horas("5:30") == 5.5
    assert _horas("-05:10") == -5.17 and _horas("-0:30") == -0.5 and _horas("+01:00") == 1.0
    assert _horas(time(5, 30)) == 5.5
    assert _horas(timedelta(hours=5, minutes=10)) == 5.17 and _horas(pd.Timedelta(hours=-2, minutes=-30)) == -2.5
    assert _horas(datetime(1900, 1, 1, 12, 0)) == 60.0  # celula [h]:mm do Excel passa de 24h
    assert _horas(None) == 0.0 and _horas(float("nan")) == 0.0 and _horas("") == 0.0


def test_converter_horas_invalido_aponta_arquivo_e_linha():
    for ruim in ("abc", "5:75", "1:2:3:4", True, "--3"):
        try:
            _horas(ruim)
        except ArquivoInvalidoError as e:
            assert "arquivo teste, linha 2" in str(e) and "saldo" in str(e)
        else:
            raise AssertionError(f"deveria recusar {ruim!r}")


def test_converter_valor_monetario():
    assert converter_valor_monetario("R$ 62,50", "x", "c") == 62.5
    assert converter_valor_monetario("1.234,56", "x", "c") == 1234.56
    assert converter_valor_monetario(None, "x", "c") is None


def _cadastro_para_bancos() -> pd.DataFrame:
    return carregar_cadastro_colaboradores(_cadastro_em_arquivo([
        _linha_cadastro(PIS_BANCO, "100", "SIDNEY TESTE", banco_horas="sim"),
        _linha_cadastro(PIS_COLEGA, "101", "COLEGA TESTE"),
        _linha_cadastro(PIS_SEM_PONTO, "102", "ADRIANO TESTE", banco_horas="sim", registra_ponto="nao"),
    ]))


def test_export_secullum_com_hhmm_negativo_virgula_e_celula_de_hora():
    cadastro = _cadastro_para_bancos()
    caminho = _pasta() / "banco_secullum.xlsx"
    pd.DataFrame({
        "matricula": ["100", "101", "102", "100"],
        "saldo": ["-05:10", time(5, 30), "1,5", 2.25],
    }).to_excel(caminho, index=False)
    out = ler_banco_horas_secullum(str(caminho), cadastro=cadastro)
    assert out["saldo_banco_horas"].tolist() == [-5.17, 5.5, 1.5, 2.25]
    assert out["pis"].tolist() == [PIS_BANCO, PIS_COLEGA, PIS_SEM_PONTO, PIS_BANCO]

    pd.DataFrame({"matricula": ["100", "101"], "saldo": ["1,5", "xx"]}).to_excel(caminho, index=False)
    try:
        ler_banco_horas_secullum(str(caminho), cadastro=cadastro)
    except ArquivoInvalidoError as e:
        assert "banco_secullum.xlsx" in str(e) and "linha 3" in str(e)
    else:
        raise AssertionError("saldo invalido deveria falhar")


def test_planilha_adriano_aceita_hhmm_e_celulas_de_hora():
    cadastro = _cadastro_para_bancos()
    caminho = _pasta() / "adriano.xlsx"
    pd.DataFrame({
        "semana": [1, 2],
        "horas_extras": ["02:00", "1,5"],
        "horas_debito": [time(0, 30), "00:15"],
        "saldo_anterior": ["-01:00", None],
    }).to_excel(caminho, index=False)
    out = ler_banco_horas_adriano(str(caminho), "102", cadastro=cadastro)
    assert out["saldo_banco_horas"].iloc[0] == round(-1 + (2 + 1.5) - (0.5 + 0.25), 2) == 1.75

    pd.DataFrame({"semana": [1, 2], "horas_extras": [1.0, 2.0], "horas_debito": [0.0, "xx"]}).to_excel(caminho, index=False)
    try:
        ler_banco_horas_adriano(str(caminho), "102", cadastro=cadastro)
    except ArquivoInvalidoError as e:
        assert "adriano.xlsx" in str(e) and "linha 3" in str(e) and "horas_debito" in str(e)
    else:
        raise AssertionError("horas invalidas deveriam falhar")


# ---------------------------------------------------------------------------
# Preflight: banco x cadastro
# ---------------------------------------------------------------------------
def _banco(pis: str, saldo: float, fonte: str = "secullum_sidney") -> pd.DataFrame:
    return pd.DataFrame([{"pis": pis, "saldo_banco_horas": saldo, "fonte": fonte}])


def test_preflight_recusa_saldo_de_quem_nao_esta_no_banco():
    cadastro = _cadastro_para_bancos()
    executar_checagens_preflight(cadastro, None, _banco(PIS_BANCO, 5.0), _banco(PIS_SEM_PONTO, 2.0, "planilha_adriano"))
    for kwargs in ({"df_banco_secullum": _banco(PIS_COLEGA, 3.0)},
                   {"df_banco_adriano": _banco(PIS_COLEGA, 3.0, "planilha_adriano")}):
        try:
            executar_checagens_preflight(cadastro, **kwargs, caminho_banco_secullum="banco.xlsx", caminho_banco_adriano="ad.xlsx")
        except ArquivoInvalidoError as e:
            msg = str(e)
            assert "COLEGA TESTE" in msg and "banco_horas" in msg
            assert "pagas na folha" in msg and "contadas no banco de horas" in msg
        else:
            raise AssertionError("deveria recusar saldo de quem nao e' banco de horas")


def test_preflight_aceita_linha_de_saldo_zero_de_quem_nao_esta_no_banco():
    # Export do Secullum que lista todo mundo: saldo 0 nao paga nada duas vezes.
    executar_checagens_preflight(_cadastro_para_bancos(), None, _banco(PIS_COLEGA, 0.0))


def test_avisos_de_banco_sem_saldo_nos_arquivos():
    cadastro = _cadastro_para_bancos()
    avisos = avisos_banco_horas(cadastro, _banco(PIS_BANCO, 5.0), None)
    assert len(avisos) == 1 and "ADRIANO TESTE" in avisos[0] and "0999" in avisos[0]
    assert len(avisos_banco_horas(cadastro, None, None)) == 2
    assert avisos_banco_horas(cadastro, _banco(PIS_BANCO, 5.0), _banco(PIS_SEM_PONTO, 1.0, "planilha_adriano")) == []
    assert avisos_banco_horas(cadastro.drop(columns=["banco_horas"]), None, None) == []  # cadastro antigo


# ---------------------------------------------------------------------------
# Relacao de valores: roteamento para o banco, sem ponto, DSR
# ---------------------------------------------------------------------------
def _afd(pis: str, nsr0: int, dias: dict[date, list[tuple[int, int]]]) -> list[str]:
    linhas = []
    for dia, horarios in dias.items():
        for h, m in horarios:
            linhas.append(f"{nsr0 + len(linhas):09d}3{pis.zfill(12)}{datetime(dia.year, dia.month, dia.day, h, m):%Y%m%d%H%M}")
    return linhas


def _cenario(extras_banco: bool = True):
    """
    Setembro/2026. Banco e Colega fazem o MESMO ponto: 01/09 +2h (HE 50), 02/09 1h de atraso,
    domingo 06/09 4h (HE 100) e faltam em 03/09 (sem batida) e 08/09 (2 semanas ISO diferentes
    -> 2 DSR perdidos, 2 dias de falta). O Sem-ponto nao tem nenhuma batida, de proposito.
    """
    cadastro = _cadastro_para_bancos()
    normal = [(8, 0), (12, 0), (13, 0), (17, 0)]
    dias = {}
    for dia in pd.date_range("2026-09-01", "2026-09-30").date:
        if dia.weekday() < 5 and dia not in (date(2026, 9, 3), date(2026, 9, 8)):
            dias[dia] = normal
    dias[date(2026, 9, 1)] = [(8, 0), (12, 0), (13, 0), (19, 0)]
    dias[date(2026, 9, 2)] = [(8, 0), (12, 0), (13, 0), (16, 0)]
    dias[date(2026, 9, 6)] = [(8, 0), (12, 0)]
    linhas = _afd(PIS_BANCO, 1, dias)
    linhas += _afd(PIS_COLEGA, len(linhas) + 1, dias)
    caminho = _pasta() / "ponto.txt"
    caminho.write_text("\n".join(linhas) + "\n", encoding="latin-1")
    apuracao, fila = rodada_1_gerar_fila_de_excecoes(
        str(caminho), cadastro, date(2026, 9, 1), date(2026, 9, 30), feriados={}
    )
    return cadastro, apuracao, fila


def _rodar(cadastro, apuracao, saldos=None, **kwargs) -> pd.DataFrame:
    saldos = saldos if saldos is not None else {PIS_BANCO: 12.5, PIS_SEM_PONTO: 7.0}
    banco = pd.DataFrame({"pis": cadastro["pis"], "saldo_banco_horas": cadastro["pis"].map(saldos).fillna(0.0)})
    consignados = pd.DataFrame({"pis": cadastro["pis"], "valor_consignado": 0.0})
    return rodada_2_gerar_relacao_de_valores(apuracao, cadastro, consignados, banco, None, "09/2026", **kwargs)


def _linha(relacao: pd.DataFrame, pis_matricula: str) -> pd.Series:
    return relacao[relacao["codigo_folha"].astype(str) == pis_matricula].iloc[0]


def test_sem_ponto_fica_fora_da_apuracao():
    cadastro, apuracao, fila = _cenario()
    assert PIS_SEM_PONTO not in set(apuracao["pis"]) and PIS_SEM_PONTO not in set(fila["pis"])
    assert {PIS_BANCO, PIS_COLEGA} == set(apuracao["pis"])


def test_quem_esta_no_banco_nao_recebe_hora_do_relogio_na_folha():
    cadastro, apuracao, _ = _cenario()
    rel = _rodar(cadastro, apuracao)
    banco, colega = _linha(rel, "100"), _linha(rel, "101")

    # colega (folha normal) mantem tudo o que o relogio apurou
    assert (colega["0150"], colega["0200"], colega["8069"]) == (2.0, 4.0, 1.0)
    assert colega["_banco_horas"] == "Não" and colega["_he50_banco"] == colega["_he100_banco"] == 0.0
    # banco: nada vai para a folha, e o que foi desviado aparece nas colunas auxiliares
    assert (banco["0150"], banco["0200"], banco["8069"]) == (0.0, 0.0, 0.0)
    assert banco["_banco_horas"] == "Sim"
    assert (banco["_he50_banco"], banco["_he100_banco"], banco["_atraso_falta_parcial_banco"]) == (2.0, 4.0, 1.0)
    assert banco["0999"] == 12.5  # o saldo continua vindo da fonte existente
    # falta de dia inteiro NAO e' afetada
    assert banco["8792"] == colega["8792"] == 2
    assert banco["_dias_falta_injustificada"] == 2 and banco["_dias_desconto_va"] == 2
    assert banco["_perde_premio_copr"] is True


def test_constantes_de_premissa_controlam_he100_e_atraso():
    cadastro, apuracao, _ = _cenario()
    antigo = (pipeline.BANCO_ABSORVE_HE_100, pipeline.BANCO_ABSORVE_ATRASO_E_FALTA_PARCIAL)
    try:
        pipeline.BANCO_ABSORVE_HE_100 = False
        pipeline.BANCO_ABSORVE_ATRASO_E_FALTA_PARCIAL = False
        banco = _linha(_rodar(cadastro, apuracao), "100")
    finally:
        pipeline.BANCO_ABSORVE_HE_100, pipeline.BANCO_ABSORVE_ATRASO_E_FALTA_PARCIAL = antigo
    assert (banco["0150"], banco["0200"], banco["8069"]) == (0.0, 4.0, 1.0)
    assert banco["_he100_banco"] == 0.0 and banco["_atraso_falta_parcial_banco"] == 0.0


def test_sem_ponto_zera_codigos_do_relogio_e_copr_vira_conferir():
    cadastro, apuracao, _ = _cenario()
    sem = _linha(_rodar(cadastro, apuracao), "102")
    assert (sem["0150"], sem["0200"], sem["8069"], sem["8792"], sem["8794"]) == (0, 0, 0, 0, 0)
    assert sem["_perde_premio_copr"] == "Conferir"
    assert sem["0999"] == 7.0 and sem["_dias_falta_total"] == 0
    assert _linha(_rodar(cadastro, apuracao), "101")["_perde_premio_copr"] is True   # quem bate ponto segue Perde/Nao Perde


def test_conferencia_secullum_compara_o_valor_do_relogio_e_nao_o_roteado():
    cadastro, apuracao, _ = _cenario()
    apuracao.attrs["totais_secullum"] = pd.DataFrame(
        [{"pis": PIS_BANCO, "faltas_dias": 2.0, "faltas_horas": 60, "extra_50": 120, "extra_100": 240}]
    )
    banco = _linha(_rodar(cadastro, apuracao), "100")
    assert banco["0150"] == 0.0                       # roteado ao banco...
    assert banco["_sec_he50"] == 2.0 and banco["_dif_he50"] == 0.0     # ...mas a conferencia usa o do relogio
    assert banco["_dif_he100"] == 0.0 and banco["_dif_horas_falta"] == 0.0


def test_cadastro_sem_as_colunas_novas_da_o_mesmo_resultado_de_antes():
    cadastro, apuracao, _ = _cenario()
    antigo = cadastro.drop(columns=["banco_horas", "registra_ponto"])
    rel = _rodar(antigo, apuracao, saldos={})
    for matricula in ("100", "101"):
        linha = _linha(rel, matricula)
        assert (linha["0150"], linha["0200"], linha["8069"], linha["8792"]) == (2.0, 4.0, 1.0, 2)
        assert linha["_banco_horas"] == "Não"
        assert linha["_perde_premio_copr"] is True


def test_exportacao_mostra_colunas_auxiliares_e_nao_destaca_conferir():
    cadastro, apuracao, _ = _cenario()
    arquivos = exportar_por_empresa(_rodar(cadastro, apuracao), str(_pasta()))
    from openpyxl import load_workbook

    wb = load_workbook(arquivos[0])
    assert wb["Contabilidade"].max_column == 12   # layout da contabilidade intacto
    ws = wb["Conferência RH"]
    cabecalhos = [c.value for c in ws[1]]
    for esperado in ("Banco de horas", "HE 50% enviada ao banco", "HE 100% enviada ao banco",
                     "Atraso/falta parcial enviados ao banco"):
        assert esperado in cabecalhos, esperado
    assert "Valor do dia de DSR" not in cabecalhos    # so' no modo "valor"
    col_copr = cabecalhos.index("Prêmio COPR") + 1
    linhas = {str(ws.cell(row=i, column=4).value): i for i in range(2, ws.max_row + 1)}
    sem_ponto = ws.cell(row=linhas["102"], column=col_copr)
    assert sem_ponto.value == "Conferir" and not str(sem_ponto.fill.fgColor.rgb).endswith("FBE4E1")
    perde = ws.cell(row=linhas["100"], column=col_copr)
    assert perde.value == "Perde" and str(perde.fill.fgColor.rgb).endswith("FBE4E1")


# ---------------------------------------------------------------------------
# DSR (8794)
# ---------------------------------------------------------------------------
def test_dsr_em_dias_e_o_numero_de_semanas_perdidas():
    cadastro, apuracao, _ = _cenario()
    rel = _rodar(cadastro, apuracao)    # padrao
    assert pipeline.UNIDADE_DSR_PADRAO == "dias"
    for matricula in ("100", "101"):
        linha = _linha(rel, matricula)
        assert linha["8794"] == linha["_semanas_perde_dsr"] == 2
    assert _linha(rel, "102")["8794"] == 0
    assert "_valor_dia_dsr" not in rel.columns


def test_dsr_em_valor_multiplica_semanas_pelo_valor_do_dia():
    cadastro, apuracao, _ = _cenario()
    cadastro = cadastro.assign(valor_dia_dsr=[62.5, 70.0, float("nan")])   # o 3o nao perde DSR: sem valor e' ok
    rel = _rodar(cadastro, apuracao, unidade_dsr="valor")
    assert _linha(rel, "100")["8794"] == 125.0 and _linha(rel, "100")["_valor_dia_dsr"] == 62.5
    assert _linha(rel, "101")["8794"] == 140.0
    assert _linha(rel, "102")["8794"] == 0
    # o dict antigo continua valendo e sobrepoe a coluna
    rel = _rodar(cadastro, apuracao, unidade_dsr="valor", valor_dia_dsr_por_pis={PIS_BANCO: 100.0})
    assert _linha(rel, "100")["8794"] == 200.0 and _linha(rel, "101")["8794"] == 140.0
    # o cabecalho auxiliar aparece na Conferencia RH
    from openpyxl import load_workbook
    wb = load_workbook(exportar_por_empresa(rel, str(_pasta()))[0])
    assert "Valor do dia de DSR" in [c.value for c in wb["Conferência RH"][1]]


def test_dsr_em_valor_sem_valor_para_quem_perdeu_da_erro_e_nunca_zero():
    cadastro, apuracao, _ = _cenario()
    for tentativa in (cadastro, cadastro.assign(valor_dia_dsr=[62.5, float("nan"), float("nan")])):
        try:
            _rodar(tentativa, apuracao, unidade_dsr="valor")
        except ArquivoInvalidoError as e:
            msg = str(e)
            assert "COLEGA TESTE" in msg and "valor_dia_dsr" in msg
            assert "ADRIANO TESTE" not in msg   # quem nao perdeu DSR nao precisa de valor
        else:
            raise AssertionError("deveria exigir o valor do dia de DSR")


def test_unidade_dsr_desconhecida_e_recusada():
    cadastro, apuracao, _ = _cenario()
    try:
        _rodar(cadastro, apuracao, unidade_dsr="reais")
    except ValueError as e:
        assert "dias" in str(e) and "valor" in str(e)
    else:
        raise AssertionError("unidade invalida deveria falhar")


# ---------------------------------------------------------------------------
# Cenario real de 09/2026 (so' roda se a pasta existir)
# ---------------------------------------------------------------------------
PASTA_REAL = Path("Dados e Docs Reais")
MATRICULA_BANCO_REAL = "100"   # SIDNEY DELFINO MAIOL - NAO confirmado com a cliente (pode ser o Sidnei, matricula 80)


def _cenario_real():
    from pipeline import carregar_consignados
    from banco_horas import consolidar_banco_horas

    cadastro = carregar_cadastro_colaboradores(str(PASTA_REAL / "cadastro_sintetico.xlsx"))
    apuracao, fila = rodada_1_gerar_fila_de_excecoes(
        str(PASTA_REAL / "Cartão Ponto 09.2026.pdf"), cadastro, date(2026, 8, 26), date(2026, 9, 25),
        caminho_espelho=str(PASTA_REAL / "Espelho de Ponto 09.2026.pdf"),
    )
    consignados = carregar_consignados(str(PASTA_REAL / "Consignado Modificado.xlsx"), cadastro=cadastro)
    banco = consolidar_banco_horas(cadastro["pis"].tolist(), None, None)   # nao existe export de banco real
    return cadastro, apuracao, consignados, banco


def _antes_e_depois_real(**kwargs):
    cadastro, apuracao, consignados, banco = _cenario_real()
    antes = rodada_2_gerar_relacao_de_valores(apuracao, cadastro, consignados, banco, None, "09/2026", **kwargs)
    marcado = cadastro.assign(banco_horas=cadastro["matricula"].astype(str).str.strip() == MATRICULA_BANCO_REAL)
    depois = rodada_2_gerar_relacao_de_valores(apuracao, marcado, consignados, banco, None, "09/2026", **kwargs)
    return cadastro, antes, depois, marcado, banco


def test_cenario_real_setembro_2026_banco_e_dsr():
    if not (PASTA_REAL / "Espelho de Ponto 09.2026.pdf").exists():
        return  # pasta real ausente (fora do git): nada a verificar
    _, antes, depois, marcado, banco = _antes_e_depois_real()
    a, d = _linha(antes, MATRICULA_BANCO_REAL), _linha(depois, MATRICULA_BANCO_REAL)
    assert a["0150"] > 0 and d["0150"] == 0 and d["0200"] == 0 and d["8069"] == 0
    assert d["_he50_banco"] == a["0150"]
    outros = antes["codigo_folha"].astype(str) != MATRICULA_BANCO_REAL
    cols = ["0150", "0200", "8069", "8792", "8794", "0981", "0999"]
    assert antes.loc[outros, cols].equals(depois.loc[outros, cols])    # colegas intactos
    assert (antes["8794"] == antes["_semanas_perde_dsr"]).all()
    assert len(avisos_banco_horas(marcado, None, None)) == 1   # sem export de banco, o 0999 dele sai 0


if __name__ == "__main__":
    testes = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    for nome, funcao in testes:
        funcao()
        print(f"OK  {nome}")
    print(f"\n{len(testes)} teste(s) passaram.")

    if (PASTA_REAL / "Espelho de Ponto 09.2026.pdf").exists():
        pd.set_option("display.width", 250)
        pd.set_option("display.max_columns", 40)
        cadastro, antes, depois, marcado, banco = _antes_e_depois_real()
        print(f"\n=== Cenario real 09/2026 - banco de horas = matricula {MATRICULA_BANCO_REAL} (NAO confirmado) ===")
        a, d = _linha(antes, MATRICULA_BANCO_REAL), _linha(depois, MATRICULA_BANCO_REAL)
        print(a["nome_colaborador"])
        for codigo in ("0150", "0200", "8069", "0999"):
            print(f"  {codigo}: antes={a[codigo]}  depois={d[codigo]}")
        print("  desviado ao banco:", d["_he50_banco"], d["_he100_banco"], d["_atraso_falta_parcial_banco"])
        print("  avisos:", avisos_banco_horas(marcado, None, None))

        print("\n=== 8794 dos colaboradores ===")
        valor_ilustrativo = {p: 100.0 for p in cadastro["pis"]}   # R$ 100,00/dia SO' para ilustrar o modo "valor"
        em_valor = rodada_2_gerar_relacao_de_valores(
            _cenario_real()[1], cadastro, _cenario_real()[2], banco, None, "09/2026",
            unidade_dsr="valor", valor_dia_dsr_por_pis=valor_ilustrativo,
        )
        tabela = antes[["codigo_folha", "nome_colaborador", "8794", "_semanas_perde_dsr"]].rename(
            columns={"8794": "8794 (dias)"}
        ).assign(**{"8794 (R$ 100/dia, ilustrativo)": em_valor["8794"].values})
        print(tabela.to_string(index=False))
        print(f"total em dias: {int(antes['8794'].sum())}")
