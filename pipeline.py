"""
pipeline.py
===========
Orquestra o fechamento de folha ponta a ponta:

    AFD (.txt) + cadastro de colaboradores + consignados (Excel)
    + banco de horas (Secullum / planilha Adriano) + decisoes do RH
    (vindas do painel de aprovacao)
        -> "Relacao de Valores" (um arquivo por empresa, no layout
           exigido pela contabilidade)
        -> fila de excecoes para validacao do RH (antes do fechamento)

Como rodar (depois de ajustar os caminhos dos arquivos reais no bloco
`if __name__ == "__main__"` no final do arquivo):

    python pipeline.py

Duas rodadas por fechamento:
  1. Rodada de EXCECOES: gera `fila_validacao_rh.csv`, que e' carregada no
     dashboard_aprovacao.html para o RH/gestor classificar cada dia
     estranho (falta real, erro de relogio, atestado, etc.).
  2. Rodada FINAL: le as decisoes exportadas do dashboard
     (`decisoes_rh.csv`) e gera a Relacao de Valores definitiva.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, time
from pathlib import Path

import pandas as pd

from afd_parser import parse_afd, agrupar_batidas_por_dia
from business_rules import Jornada, apurar_periodo, consolidar_mes, gerar_fila_validacao_rh
from banco_horas import consolidar_banco_horas, ler_banco_horas_secullum, ler_banco_horas_adriano
from validacao import verificar_colunas, validar_pis_existem, normalizar_identificador_colaborador, ArquivoInvalidoError
from leitura_arquivos import ler_arquivo_generico

# ---------------------------------------------------------------------------
# Codigos de evento exigidos pela contabilidade (Secao 4 do briefing)
# ---------------------------------------------------------------------------
COD_HE_50 = "0150"
COD_HE_100 = "0200"
COD_HORAS_FALTA = "8069"
COD_DIAS_FALTA = "8792"
COD_DESCONTO_DSR = "8794"
COD_ADIANTAMENTO_CONSIGNADO = "0981"
COD_OUTROS_BANCO_HORAS = "0999"

COLUNAS_RELACAO_VALORES = [
    "codigo_empresa", "razao_social", "competencia",
    "codigo_folha", "nome_colaborador",
    COD_HE_50, COD_HE_100, COD_HORAS_FALTA, COD_DIAS_FALTA,
    COD_DESCONTO_DSR, COD_ADIANTAMENTO_CONSIGNADO, COD_OUTROS_BANCO_HORAS,
]

# Colunas que o Cadastro de Colaboradores precisa ter - validadas logo no
# carregamento, para o erro apontar exatamente o arquivo/coluna em vez de
# estourar um KeyError sem contexto no meio do calculo.
COLUNAS_OBRIGATORIAS_CADASTRO = [
    "pis", "matricula", "nome", "codigo_empresa", "razao_social",
    "entrada", "saida_almoco", "retorno_almoco", "saida", "dias_trabalho",
]


def carregar_cadastro_colaboradores(caminho: str) -> pd.DataFrame:
    """
    Cadastro mestre. Colunas esperadas (ajuste os nomes se seu cadastro
    real usar outros):
        pis, matricula, nome, codigo_empresa, razao_social,
        entrada, saida_almoco, retorno_almoco, saida, dias_trabalho
        (dias_trabalho = string "0,1,2,3,4" para seg-sex, 0=segunda)

    Levanta `ArquivoInvalidoError` (com o caminho e a(s) coluna(s) que
    faltam) se o arquivo nao tiver todas as colunas obrigatorias.
    """
    df = ler_arquivo_generico(caminho, "Cadastro de Colaboradores")
    verificar_colunas(df, COLUNAS_OBRIGATORIAS_CADASTRO, "Cadastro de Colaboradores", caminho)
    df["pis"] = df["pis"].astype(str).str.strip()
    return df


def montar_jornadas(df_cadastro: pd.DataFrame) -> dict[str, Jornada]:
    jornadas = {}
    for _, row in df_cadastro.iterrows():
        try:
            jornadas[row["pis"]] = Jornada(
                entrada=pd.to_datetime(row["entrada"]).time(),
                saida_almoco=pd.to_datetime(row["saida_almoco"]).time(),
                retorno_almoco=pd.to_datetime(row["retorno_almoco"]).time(),
                saida=pd.to_datetime(row["saida"]).time(),
                dias_trabalho=set(int(d) for d in str(row["dias_trabalho"]).split(",")),
            )
        except Exception as e:
            raise ArquivoInvalidoError(
                f"Não consegui interpretar a jornada do colaborador com PIS '{row.get('pis')}' "
                f"no Cadastro de Colaboradores. Confira se 'entrada', 'saida_almoco', "
                f"'retorno_almoco' e 'saida' estão em formato de hora (ex.: 08:00) e se "
                f"'dias_trabalho' está no formato \"0,1,2,3,4\" (0=segunda). Detalhe: {e}"
            ) from e
    return jornadas


def carregar_consignados(caminho: str, cadastro: pd.DataFrame | None = None) -> pd.DataFrame:
    """
    Extrato de consignados. Espera colunas matricula/pis + valor_desconto.

    Passe `cadastro` (o DataFrame já carregado de
    `carregar_cadastro_colaboradores`) sempre que possível: muitos
    extratos de consignados identificam o colaborador pela MATRÍCULA
    interna, não pelo PIS, e sem o cadastro em mãos não dá para saber
    qual dos dois está na coluna de identificação — veja
    `validacao.normalizar_identificador_colaborador` para o porquê.
    """
    df = ler_arquivo_generico(caminho, "Extrato de Consignados")
    col_id = next((c for c in df.columns if "pis" in c or "matricula" in c), None)
    col_valor = next((c for c in df.columns if "valor" in c or "desconto" in c), None)
    if col_id is None or col_valor is None:
        raise ArquivoInvalidoError(
            f"O Extrato de Consignados ('{caminho}') está sem uma coluna de identificação "
            f"(pis/matrícula) e/ou de valor (valor/desconto).\n"
            f"Colunas encontradas: {list(df.columns)}."
        )
    out = df[[col_id, col_valor]].rename(columns={col_id: "pis", col_valor: "valor_consignado"})
    out["pis"] = out["pis"].astype(str).str.strip()
    if cadastro is not None:
        out = normalizar_identificador_colaborador(out, "pis", cadastro, "Extrato de Consignados", caminho)
    return out.groupby("pis", as_index=False)["valor_consignado"].sum()


def executar_checagens_preflight(
    df_cadastro: pd.DataFrame,
    df_consignados: pd.DataFrame | None = None,
    df_banco_secullum: pd.DataFrame | None = None,
    df_banco_adriano: pd.DataFrame | None = None,
    caminho_consignados: str = "",
    caminho_banco_secullum: str = "",
    caminho_banco_adriano: str = "",
) -> None:
    """
    Roda ANTES do processamento pesado (leitura/apuracao do AFD): confere
    que todo PIS/matrícula citado nos arquivos opcionais existe de fato no
    Cadastro de Colaboradores. Sem essa checagem, um PIS digitado errado
    no banco de horas do Adriano, por exemplo, simplesmente nao casaria
    com ninguem no merge la' na frente - a folha de alguem sairia sem o
    banco de horas dela, silenciosamente, sem nenhum aviso.
    """
    pis_validos = df_cadastro["pis"]

    if df_consignados is not None and not df_consignados.empty:
        validar_pis_existem(df_consignados["pis"], pis_validos, "Extrato de Consignados", caminho_consignados)

    if df_banco_secullum is not None and not df_banco_secullum.empty:
        validar_pis_existem(
            df_banco_secullum["pis"], pis_validos, "Banco de Horas do Secullum", caminho_banco_secullum
        )

    if df_banco_adriano is not None and not df_banco_adriano.empty:
        validar_pis_existem(
            df_banco_adriano["pis"], pis_validos, "Planilha semanal do Adriano", caminho_banco_adriano
        )


def rodada_1_gerar_fila_de_excecoes(
    caminho_afd: str,
    df_cadastro: pd.DataFrame,
    data_inicio: date,
    data_fim: date,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Retorna (apuracao_diaria_completa, fila_para_validacao_rh)."""
    resultado_afd = parse_afd(caminho_afd)
    if resultado_afd.erros_leitura:
        print(f"[aviso] {len(resultado_afd.erros_leitura)} linha(s) do AFD nao puderam ser lidas - ver log.")

    batidas_por_dia = agrupar_batidas_por_dia(resultado_afd.batidas)
    jornadas = montar_jornadas(df_cadastro)
    apuracao = apurar_periodo(batidas_por_dia, jornadas, data_inicio, data_fim)
    fila = gerar_fila_validacao_rh(apuracao)
    return apuracao, fila


def exportar_fila_validacao_rh(fila: pd.DataFrame, df_cadastro: pd.DataFrame, caminho_csv: str) -> pd.DataFrame:
    """
    Prepara e grava a fila de excecoes num CSV pronto para o RH abrir
    direto no Excel:
      - cruza com o cadastro pra trazer nome_colaborador e razao_social
        (o RH nao precisa abrir outro arquivo pra saber quem e' quem);
      - ordena por nome do colaborador e, dentro do mesmo colaborador,
        cronologicamente pela data (a ordenacao cronologica acontece
        ANTES de converter a data pra string, senao "05/12/2026" ficaria
        antes de "20/01/2026" numa ordenacao de texto);
      - formata a data como dd/mm/aaaa (padrao brasileiro);
      - grava em 'utf-8-sig', que e' o utf-8 com o BOM que faz o Excel
        abrir acentos corretamente sem pedir importacao de texto.
    Retorna o DataFrame final tambem, para quem quiser inspecionar sem
    reabrir o CSV.
    """
    if fila.empty:
        fila_final = fila.assign(nome_colaborador=pd.Series(dtype=str), razao_social=pd.Series(dtype=str))
        fila_final.to_csv(caminho_csv, index=False, encoding="utf-8-sig")
        return fila_final

    info_colaborador = df_cadastro[["pis", "nome", "razao_social"]].rename(columns={"nome": "nome_colaborador"})
    fila_com_nome = fila.merge(info_colaborador, on="pis", how="left")

    fila_com_nome["data"] = pd.to_datetime(fila_com_nome["data"])
    fila_ordenada = fila_com_nome.sort_values(["nome_colaborador", "data"]).reset_index(drop=True)
    fila_ordenada["data"] = fila_ordenada["data"].dt.strftime("%d/%m/%Y")

    fila_final = fila_ordenada[["pis", "nome_colaborador", "razao_social", "data", "status", "motivo_validacao"]]
    fila_final.to_csv(caminho_csv, index=False, encoding="utf-8-sig")
    return fila_final


def rodada_2_gerar_relacao_de_valores(
    apuracao_diaria: pd.DataFrame,
    df_cadastro: pd.DataFrame,
    df_consignados: pd.DataFrame,
    df_banco_horas: pd.DataFrame,
    decisoes_rh: pd.DataFrame | None,
    competencia: str,
    valor_dia_dsr_por_pis: dict[str, float] | None = None,
) -> pd.DataFrame:
    """Gera a tabela final no layout exigido pela contabilidade, uma linha por colaborador."""
    resumo_mes = consolidar_mes(apuracao_diaria, decisoes_rh)

    df = df_cadastro.merge(resumo_mes, on="pis", how="left")
    df = df.merge(df_consignados, on="pis", how="left")
    df = df.merge(df_banco_horas, on="pis", how="left")

    for col in ["horas_extra_50", "horas_extra_100", "dias_falta", "dias_desconto_va",
                "semanas_perde_dsr", "valor_consignado", "saldo_banco_horas"]:
        if col in df.columns:
            df[col] = df[col].fillna(0)

    valor_dia_dsr_por_pis = valor_dia_dsr_por_pis or {}
    df["valor_desconto_dsr"] = df.apply(
        lambda r: round(r["semanas_perde_dsr"] * valor_dia_dsr_por_pis.get(r["pis"], 0.0), 2), axis=1
    )

    saida = pd.DataFrame({
        "codigo_empresa": df["codigo_empresa"],
        "razao_social": df["razao_social"],
        "competencia": competencia,
        "codigo_folha": df["matricula"],
        "nome_colaborador": df["nome"],
        COD_HE_50: df["horas_extra_50"],
        COD_HE_100: df["horas_extra_100"],
        COD_HORAS_FALTA: df["dias_falta"] * 0,  # preencher com horas de falta reais se a jornada variar por dia
        COD_DIAS_FALTA: df["dias_falta"],
        COD_DESCONTO_DSR: df["valor_desconto_dsr"],
        COD_ADIANTAMENTO_CONSIGNADO: df["valor_consignado"],
        COD_OUTROS_BANCO_HORAS: df["saldo_banco_horas"],
        # colunas auxiliares - uteis para conferencia interna do RH, NAO
        # fazem parte do layout formal enviado a contabilidade. Ficam com
        # prefixo "_" de proposito, para o exportador saber separa-las
        # automaticamente na aba "Conferência RH" e nunca vazarem para a
        # aba "Contabilidade":
        "_perde_premio_copr": df["perde_copr"].fillna(False),
        "_dias_falta_total": df["dias_falta"],
        "_dias_desconto_va": df["dias_desconto_va"],
        "_semanas_perde_dsr": df["semanas_perde_dsr"],
    })
    return saida


def exportar_por_empresa(df_relacao: pd.DataFrame, pasta_saida: str) -> list[Path]:
    """
    Gera um .xlsx por empresa com DUAS abas:
      - "Contabilidade": estritamente as colunas do layout exigido pelo
        escritorio contabil (nenhuma coluna auxiliar, para nao arriscar
        rejeicao no sistema deles).
      - "Conferência RH": as mesmas colunas + as colunas auxiliares
        (prefixo "_" em `df_relacao`) para auditoria interna, com
        destaque visual de quem perde premio/tem faltas.
    """
    pasta = Path(pasta_saida)
    pasta.mkdir(parents=True, exist_ok=True)
    arquivos = []
    for (codigo, razao), grupo in df_relacao.groupby(["codigo_empresa", "razao_social"]):
        nome_arquivo = f"relacao_de_valores_{codigo}_{razao.replace(' ', '_')}.xlsx"
        caminho = pasta / nome_arquivo
        _escrever_workbook_formatado(grupo, caminho)
        arquivos.append(caminho)
    return arquivos


# ---------------------------------------------------------------------------
# Exportação Excel com openpyxl - layout, formatação e formatação condicional
# ---------------------------------------------------------------------------

_COR_CABECALHO = "1B4D3E"       # verde corporativo
_COR_TEXTO_CABECALHO = "FFFFFF"
_COR_BORDA = "D9D9D9"           # cinza claro
_COR_ALERTA_FUNDO = "FBE4E1"    # vermelho bem suave
_COR_ALERTA_TEXTO = "8C2E22"
LIMITE_FALTAS_ALTO = 3          # a partir de quantas faltas o RH ve' o destaque de "elevado"

# Nomes de exibicao das colunas do layout contabil (na ordem em que aparecem)
_COLUNAS_CONTABIL = [
    "codigo_empresa", "razao_social", "competencia", "codigo_folha", "nome_colaborador",
    COD_HE_50, COD_HE_100, COD_HORAS_FALTA, COD_DIAS_FALTA,
    COD_DESCONTO_DSR, COD_ADIANTAMENTO_CONSIGNADO, COD_OUTROS_BANCO_HORAS,
]
_CABECALHOS_CONTABIL = {
    "codigo_empresa": "Código Empresa", "razao_social": "Razão Social", "competencia": "Competência",
    "codigo_folha": "Código Folha", "nome_colaborador": "Nome do Colaborador",
    COD_HE_50: "HE 50% (0150)", COD_HE_100: "HE 100% (0200)",
    COD_HORAS_FALTA: "Horas Falta (8069)", COD_DIAS_FALTA: "Dias Falta (8792)",
    COD_DESCONTO_DSR: "Desconto DSR (8794)", COD_ADIANTAMENTO_CONSIGNADO: "Consignado (0981)",
    COD_OUTROS_BANCO_HORAS: "Banco de Horas (0999)",
}
_COLUNAS_AUX = ["_perde_premio_copr", "_dias_falta_total", "_dias_desconto_va", "_semanas_perde_dsr"]
_CABECALHOS_AUX = {
    "_perde_premio_copr": "Prêmio COPR", "_dias_falta_total": "Dias de Falta (total)",
    "_dias_desconto_va": "Dias Desconto VA", "_semanas_perde_dsr": "Semanas c/ Perda de DSR",
}

# Colunas monetarias/numericas que levam formato "0,00" (2 casas decimais explicitas)
_COLUNAS_DECIMAIS = {COD_HE_50, COD_HE_100, COD_HORAS_FALTA, COD_DESCONTO_DSR, COD_ADIANTAMENTO_CONSIGNADO, COD_OUTROS_BANCO_HORAS}
# Colunas de quantidade inteira (ainda assim formatadas com precisao explicita)
_COLUNAS_INTEIRAS = {COD_DIAS_FALTA, "_dias_falta_total", "_dias_desconto_va", "_semanas_perde_dsr"}
# Colunas centralizadas (codigos, matricula, competencia)
_COLUNAS_CENTRALIZADAS = {"codigo_empresa", "competencia", "codigo_folha", "_perde_premio_copr"}


def _escrever_workbook_formatado(grupo: pd.DataFrame, caminho: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    fonte_padrao = "Arial"
    fill_cabecalho = PatternFill("solid", fgColor=_COR_CABECALHO)
    fonte_cabecalho = Font(name=fonte_padrao, bold=True, color=_COR_TEXTO_CABECALHO)
    borda_fina = Border(*[Side(style="thin", color=_COR_BORDA)] * 4)
    fill_alerta = PatternFill("solid", fgColor=_COR_ALERTA_FUNDO)
    fonte_alerta = Font(name=fonte_padrao, color=_COR_ALERTA_TEXTO)

    def alinhamento_para(nome_coluna: str) -> Alignment:
        if nome_coluna in _COLUNAS_CENTRALIZADAS:
            return Alignment(horizontal="center", vertical="center")
        if nome_coluna in _COLUNAS_DECIMAIS or nome_coluna in _COLUNAS_INTEIRAS:
            return Alignment(horizontal="right", vertical="center")
        return Alignment(horizontal="left", vertical="center")

    def formato_numero_para(nome_coluna: str) -> str | None:
        if nome_coluna in _COLUNAS_DECIMAIS:
            return "0.00"
        if nome_coluna in _COLUNAS_INTEIRAS:
            return "0"
        return None

    def escrever_aba(ws, colunas: list[str], cabecalhos: dict, com_alerta: bool):
        ws.append([cabecalhos.get(c, c) for c in colunas])
        for cel in ws[1]:
            cel.fill = fill_cabecalho
            cel.font = fonte_cabecalho
            cel.alignment = Alignment(horizontal="center", vertical="center")
            cel.border = borda_fina
        ws.freeze_panes = "A2"

        for _, linha in grupo.iterrows():
            valores = []
            for c in colunas:
                v = linha[c]
                if c == "_perde_premio_copr":
                    v = "Perde" if bool(v) else "Não Perde"
                valores.append(v)
            ws.append(valores)

        n_linhas = len(grupo)
        for i in range(2, n_linhas + 2):
            linha_df = grupo.iloc[i - 2]
            destacar = com_alerta and (
                bool(linha_df.get("_perde_premio_copr", False))
                or int(linha_df.get("_dias_falta_total", 0) or 0) >= LIMITE_FALTAS_ALTO
            )
            for j, c in enumerate(colunas, start=1):
                cel = ws.cell(row=i, column=j)
                cel.font = fonte_alerta if destacar else Font(name=fonte_padrao)
                cel.alignment = alinhamento_para(c)
                cel.border = borda_fina
                fmt = formato_numero_para(c)
                if fmt:
                    cel.number_format = fmt
                if destacar:
                    cel.fill = fill_alerta

        for j, c in enumerate(colunas, start=1):
            letra = get_column_letter(j)
            maior = max(
                [len(str(cabecalhos.get(c, c)))]
                + [len(str(cabecalhos.get(c, c)) if pd.isna(linha[c]) else str(linha[c])) for _, linha in grupo.iterrows()]
            )
            ws.column_dimensions[letra].width = min(max(maior + 3, 10), 42)

    wb = Workbook()
    ws_contabil = wb.active
    ws_contabil.title = "Contabilidade"
    escrever_aba(ws_contabil, _COLUNAS_CONTABIL, _CABECALHOS_CONTABIL, com_alerta=False)

    ws_rh = wb.create_sheet("Conferência RH")
    colunas_rh = _COLUNAS_CONTABIL + _COLUNAS_AUX
    cabecalhos_rh = {**_CABECALHOS_CONTABIL, **_CABECALHOS_AUX}
    escrever_aba(ws_rh, colunas_rh, cabecalhos_rh, com_alerta=True)

    wb.save(caminho)


if __name__ == "__main__":
    # --------------------------------------------------------------------
    # AJUSTE OS CAMINHOS ABAIXO PARA OS ARQUIVOS REAIS ANTES DE RODAR
    # --------------------------------------------------------------------
    CAMINHO_AFD = "dados/ponto.txt"
    CAMINHO_CADASTRO = "dados/cadastro_colaboradores.xlsx"
    CAMINHO_CONSIGNADOS = "dados/consignados.xlsx"
    CAMINHO_BANCO_HORAS_SECULLUM = "dados/banco_horas_secullum.xlsx"
    CAMINHO_PLANILHA_ADRIANO = "dados/banco_horas_adriano.xlsx"
    PIS_ADRIANO = "00000000001"
    DATA_INICIO = date(2026, 7, 26)
    DATA_FIM = date(2026, 8, 25)
    COMPETENCIA = "08/2026"

    cadastro = carregar_cadastro_colaboradores(CAMINHO_CADASTRO)

    # Le os arquivos leves primeiro (consignados, banco de horas) e roda a
    # checagem pre-voo ANTES do processamento pesado do AFD - se um PIS
    # estiver errado em algum desses arquivos, a usuaria fica sabendo em
    # segundos, sem esperar o AFD inteiro ser processado pra descobrir.
    consignados = carregar_consignados(CAMINHO_CONSIGNADOS, cadastro=cadastro)
    banco_secullum = ler_banco_horas_secullum(CAMINHO_BANCO_HORAS_SECULLUM, cadastro=cadastro) if Path(CAMINHO_BANCO_HORAS_SECULLUM).exists() else None
    banco_adriano = ler_banco_horas_adriano(CAMINHO_PLANILHA_ADRIANO, PIS_ADRIANO, cadastro=cadastro) if Path(CAMINHO_PLANILHA_ADRIANO).exists() else None

    executar_checagens_preflight(
        cadastro, consignados, banco_secullum, banco_adriano,
        caminho_consignados=CAMINHO_CONSIGNADOS,
        caminho_banco_secullum=CAMINHO_BANCO_HORAS_SECULLUM,
        caminho_banco_adriano=CAMINHO_PLANILHA_ADRIANO,
    )
    banco_horas = consolidar_banco_horas(cadastro["pis"].tolist(), banco_secullum, banco_adriano)

    # Rodada 1: gera as excecoes para o RH validar no dashboard (aqui sim
    # entra o processamento pesado - ler e apurar o AFD inteiro)
    apuracao, fila = rodada_1_gerar_fila_de_excecoes(CAMINHO_AFD, cadastro, DATA_INICIO, DATA_FIM)
    fila_exportada = exportar_fila_validacao_rh(fila, cadastro, "saida/fila_validacao_rh.csv")
    print(f"{len(fila_exportada)} excecao(oes) geradas para validacao do RH -> saida/fila_validacao_rh.csv")
    print("Abra o dashboard_aprovacao.html, carregue esse CSV, classifique as excecoes e exporte decisoes_rh.csv")

    # ---- (o RH classifica no dashboard e exporta decisoes_rh.csv aqui) ----

    # Rodada 2: com as decisoes do RH em maos, gera a Relacao de Valores final
    decisoes = pd.read_csv("saida/decisoes_rh.csv") if Path("saida/decisoes_rh.csv").exists() else None
    relacao_valores = rodada_2_gerar_relacao_de_valores(
        apuracao, cadastro, consignados, banco_horas, decisoes, COMPETENCIA
    )
    arquivos = exportar_por_empresa(relacao_valores, "saida")
    for a in arquivos:
        print(f"Gerado: {a}")
